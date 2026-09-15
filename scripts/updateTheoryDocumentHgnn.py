from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCUMENT_PATH = Path(__file__).resolve().parents[1] / "docs" / "VitaNexus_RX_Detailed_Theory_Architecture_and_Execution.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def style_table(table, widths):
    table.style = "Table Grid"
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "17395D")
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EEF5F9")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)


def move_before(element, anchor):
    anchor._p.addprevious(element._element)


def add_before(document, anchor, text="", style=None, *, bold_lead=None):
    paragraph = document.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    move_before(paragraph, anchor)
    return paragraph


document = Document(DOCUMENT_PATH)
if any(paragraph.text == "25 Baseline HGNN Specific Event Profile" for paragraph in document.paragraphs):
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for paragraph in document.paragraphs:
        if paragraph.text == "The Python FastAPI service loads the persisted feature builder, calibrated LightGBM model, twenty bootstrap models and conformal metadata. Express communicates with this service over an internal HTTP boundary. The response is accepted only when it satisfies the expected versioned schema.":
            paragraph.text = "The Python FastAPI service loads LightGBM and the protected baseline HGNN independently. LightGBM uses the persisted feature builder, calibrated selected model, twenty bootstrap models and conformal metadata. HGNN uses the hash-verified epoch 20 selection checkpoint, ordered 100-event vocabulary and saved development associations. Express communicates with both model paths over an internal HTTP boundary and accepts a response only when it satisfies the corresponding versioned schema."
        elif paragraph.text == "Clinical Safety, Adverse Risk Assessment and Recommendations operate as one visit workflow. The headings provide direct navigation to available stages. Follow-up remains locked until the required earlier stage is complete. The current stage is visually distinct and cannot trigger redundant navigation.":
            paragraph.text = "Clinical Safety, Adverse Risk Assessment, Event Profile and Recommendations operate as one visit workflow. The headings provide direct navigation to available stages. Follow-up remains locked until the required earlier stage is complete. The current stage is visually distinct and cannot trigger redundant navigation."
        elif paragraph.text == "25 Baseline HGNN Specific Event Profile":
            paragraph.paragraph_format.page_break_before = False
        elif paragraph.text == "The FAERS estimate is conditional on an adverse-event report already existing. It is not the probability that a particular patient will experience a side effect or require hospitalization. The system uses isotonic calibration, twenty case-level bootstrap models and split-conformal classification to communicate model behavior and uncertainty. Alternative medicines pass deterministic safety gates before eligible candidates are ranked using a transparent research-configured formula. PostgreSQL stores patients, consultations, analyses, model outputs, recommendations, notes, follow-ups and audit events. VitaNexus-RX is a major-project research prototype, not an autonomous prescriber or validated medical device.":
            paragraph.text = "The FAERS estimate is conditional on an adverse-event report already existing. It is not the probability that a particular patient will experience a side effect or require hospitalization. The system uses isotonic calibration, twenty case-level bootstrap models and split-conformal classification to communicate model behavior and uncertainty. A separate protected baseline HGNN ranks specific event-label associations as audit-pending supplementary evidence; its scores are not patient-incidence probabilities and never affect recommendations. Alternative medicines pass deterministic safety gates before eligible candidates are ranked using a transparent research-configured formula. PostgreSQL stores patients, consultations, analyses, model outputs, recommendations, notes, follow-ups and audit events. VitaNexus-RX is a major-project research prototype, not an autonomous prescriber or validated medical device."

        if paragraph.text in {"9.10 Notes and Follow Up", "25.8 Isolation and Remaining Audit"}:
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.page_break_before = False

    alternative_heading = next((paragraph for paragraph in document.paragraphs if paragraph.text == "9.8 Alternative Generation and Ranking"), None)
    if alternative_heading is not None:
        alternative_heading.text = "9.9 Alternative Generation and Ranking"
        followup_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "9.9 Notes and Follow Up")
        followup_heading.text = "9.10 Notes and Follow Up"
        add_before(document, alternative_heading, "9.8 Specific Event Profile", "Heading 2")
        add_before(document, alternative_heading, "The frontend loads a matching persisted HGNN profile or requests the protected baseline model through Express and FastAPI. The result contains ordered event-label model scores, coverage, versions and audit state. It remains supplementary and never enters the recommendation path.")

    output_anchor = next((paragraph for paragraph in document.paragraphs if paragraph.text == "Evidence-completeness and failure status"), None)
    output_text = "Ranked baseline HGNN event-label model scores with coverage and provenance"
    if output_anchor is not None and not any(paragraph.text == output_text for paragraph in document.paragraphs):
        add_before(document, output_anchor, output_text, "List Bullet")

    for table in document.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        first_column = [row.cells[0].text.strip() for row in table.rows]
        if header and header[0] == "Sections 1 to 9":
            table.rows[0].cells[2].text = "Sections 19 to 27"
            for row in table.rows[1:]:
                value = row.cells[2].text.strip()
                if value == "25 Conclusion":
                    row.cells[2].text = "25 Baseline HGNN Specific Event Profile"
                elif value == "26 Appendices":
                    row.cells[2].text = "26 Conclusion"
            if not any(row.cells[2].text.strip() == "27 Appendices" for row in table.rows):
                row = table.add_row()
                row.cells[2].text = "27 Appendices"
            style_table(table, [2.17, 2.17, 2.17])
            for row in table.rows:
                for cell in row.cells:
                    set_cell_margins(cell, top=55, start=120, bottom=55, end=120)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1.0
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5)
        elif header == ["Module", "Responsibility"] and "FAERS Assessment" in first_column:
            if "Specific Event Profile" not in first_column:
                row = table.add_row()
                row.cells[0].text = "Specific Event Profile"
                row.cells[1].text = "Protected baseline HGNN event-label scores, coverage, provenance and explicit failure states"
            style_table(table, [3.25, 3.25])
        elif header == ["Entity", "Purpose"] and "AdrPrediction" in first_column:
            if "HgnnEventPrediction" not in first_column:
                row = table.add_row()
                row.cells[0].text = "HgnnEventPrediction"
                row.cells[1].text = "Versioned baseline HGNN event profile"
            style_table(table, [3.25, 3.25])
        elif header == ["Area", "Principal Endpoints"] and "Event Risk Status" in first_column:
            for row in table.rows:
                if row.cells[0].text.strip() == "Event Risk Status":
                    row.cells[0].text = "Specific Event Profile"
                    row.cells[1].text = "GET and POST adverse-event-risks"
                elif row.cells[0].text.strip() == "Python Service":
                    row.cells[1].text = "GET health and HGNN status; POST LightGBM predict, batch predict and HGNN predict events"
            style_table(table, [3.25, 3.25])
        elif header == ["Technology", "Purpose"] and "PyTorch and PyTorch Geometric" in first_column:
            for row in table.rows:
                if row.cells[0].text.strip() == "PyTorch and PyTorch Geometric":
                    row.cells[1].text = "Protected baseline HGNN graph construction and inference"
            style_table(table, [3.25, 3.25])

    document.save(DOCUMENT_PATH)
    print(DOCUMENT_PATH)
    raise SystemExit(0)

replacements = {
    "The individual adverse-event page currently reports validation pending. It does not load HGNN weights, fabricate event probabilities or influence alternative ranking. The page exists as a transparent placeholder for a future validated event-level module.":
        "The Specific Event Profile loads the protected epoch 20 baseline HGNN through the internal model service. It displays ordered event-label model scores as supplementary evidence and labels the output Baseline and Audit Pending. These scores are not patient-incidence probabilities. The module never influences LightGBM, clinical safety gates, candidate generation, recommendation scoring or tie-breaking.",
    "A clinician owns patients and consultations. A patient owns conditions, allergies, medications and consultations. Each consultation can have analyses, ADR predictions, recommendation sets, notes and follow-ups. Cascade deletion applies to child analytical records, while patient deletion through the API is implemented as a soft archive.":
        "A clinician owns patients and consultations. A patient owns conditions, allergies, medications and consultations. Each consultation can have analyses, ADR predictions, versioned HGNN event profiles, recommendation sets, notes and follow-ups. Cascade deletion applies to child analytical records, while patient deletion through the API is implemented as a soft archive.",
    "An ADR prediction stores its result JSON, model version, input hash and generation time. A recommendation set stores candidate snapshots, the composite engine version and an input hash that includes indication provenance, medicines, conditions, safety evidence, model coverage and model versions. Changed inputs or artifacts therefore produce a distinct reproducible snapshot instead of silently changing the meaning of an older result.":
        "An ADR prediction stores its result JSON, model version, input hash and generation time. An HGNN event profile stores result JSON, model and checkpoint versions, model stage, audit state, input hash, coverage state and generation time. A recommendation set stores candidate snapshots, the composite engine version and an input hash that includes indication provenance, medicines, conditions, safety evidence, model coverage and model versions. Changed inputs or artifacts therefore produce a distinct reproducible snapshot instead of silently changing an older result.",
    "The experimental HGNN infrastructure is intended for future prediction of specific MedDRA adverse-event terms. It must remain separate from the active runtime until full temporal evaluation, calibration, leakage audit and clinical interpretation are completed. Event-level probabilities should remain supplementary and should not automatically influence ranking without independent validation.":
        "The protected baseline HGNN is now available through a separate supplementary runtime path. Its event-label outputs remain model scores rather than calibrated probabilities, and the interface identifies its audit as pending. Future work must complete temporal and external evaluation, calibration, subgroup analysis, leakage review and clinical interpretation before any claim of individual event probability or any proposal to change ranking.",
    "VitaNexus-RX integrates medication terminology, known interaction evidence, FAERS report classification, uncertainty information, alternative ranking and consultation persistence. DDInter identifies documented drug-drug relationships. DrugCentral supports disease checks and same-indication candidate retrieval. LightGBM classifies serious outcomes within the selected FAERS reporting context. Deterministic safety gates prevent major known risks from being overridden by the model.":
        "VitaNexus-RX integrates medication terminology, known interaction evidence, FAERS report classification, a supplementary event-label profile, uncertainty information, alternative ranking and consultation persistence. DDInter identifies documented drug-drug relationships. DrugCentral supports disease checks and same-indication candidate retrieval. LightGBM classifies serious outcomes within the selected FAERS reporting context. The protected baseline HGNN ranks event-label associations without changing any clinical safety or recommendation decision. Deterministic safety gates prevent major known risks from being overridden by either model.",
}
for paragraph in document.paragraphs:
    if paragraph.text in replacements:
        paragraph.text = replacements[paragraph.text]

conclusion = next(paragraph for paragraph in document.paragraphs if paragraph.text == "25 Conclusion")
conclusion.text = "26 Conclusion"

heading = add_before(document, conclusion, "25 Baseline HGNN Specific Event Profile", "Heading 1")
heading.paragraph_format.page_break_before = True
add_before(document, conclusion, "25.1 Purpose and Correct Interpretation", "Heading 2")
add_before(document, conclusion, "The Specific Event Profile answers a narrow supplementary question: which saved event labels receive the strongest association scores from the existing heterogeneous graph model for the encoded medicine and consultation context? A higher score means the baseline graph model learned a stronger event-label association. It does not mean that the patient has that percentage chance of developing the event. The output is not a diagnosis, incidence estimate or prescribing decision.")

add_before(document, conclusion, "25.2 Protected Model Selection", "Heading 2")
add_before(document, conclusion, "Runtime uses the hash-verified hgnn_selection_best.pt checkpoint for faers-specific-adr-hgnn-1.0.0. The post-training manifest records selection at training epoch 20, while the checkpoint stores zero-based epoch 19. Its SHA-256 is 7ca91da8449c87e65c08550083eab12cf3e267faa9e44284c496853de7fa5420. The incomplete five-epoch final-refit checkpoint remains experimental and is not served. Integration did not retrain the model, change its layers, calibrate its outputs or optimize thresholds.")

add_before(document, conclusion, "25.3 Graph Architecture", "Heading 2")
add_before(document, conclusion, "The model preserves four node types: report, drug, indication and ADR. Four lazy linear projections map their original features to 96 hidden channels. Two heterogeneous convolution layers use GraphSAGE operations, sum aggregation, residual addition and ReLU. A linear layer emits 100 logits in the exact order of the saved event vocabulary.")
architecture_rows = [
    ("Component", "Preserved implementation"),
    ("Nodes", "Report, drug, indication and ADR"),
    ("Primary context edges", "Report to primary-suspect drug, concomitant drugs and indication, with reverse relations"),
    ("Historical edges", "Drug to ADR and indication to ADR associations, with reverse relations"),
    ("Encoder", "Two HeteroConv layers with SAGEConv, sum aggregation, residual connections and ReLU"),
    ("Output", "Linear layer with 100 logits followed by one sigmoid operation at inference"),
]
table = document.add_table(rows=len(architecture_rows), cols=2)
for row_index, values in enumerate(architecture_rows):
    for column_index, value in enumerate(values):
        table.cell(row_index, column_index).text = value
style_table(table, [1.65, 4.85])
conclusion._p.addprevious(table._element)

add_before(document, conclusion, "25.4 Runtime Inputs and Graph Construction", "Heading 2")
add_before(document, conclusion, "Only age, normalized sex, candidate medicine, indication and active current medicines are accepted. Runtime reuses the training graph builder, including age scaling and missingness, sex indicators, medicine-count scaling, deterministic hashed entity features and saved development associations. Dose, route, laboratory results, allergies, genetics, renal or liver measurements and free-form notes are not invented as HGNN features.")

add_before(document, conclusion, "25.5 End to End Pipeline", "Heading 2")
pipeline = add_before(document, conclusion, "React Specific Event Profile\n  -> authenticated Express consultation route\n  -> ownership check and exact feature mapping\n  -> FastAPI HGNN status and prediction endpoints\n  -> protected checkpoint vocabulary and development associations\n  -> validated top K event-label scores\n  -> versioned PostgreSQL persistence\n  -> display only with ranking influence set to none")
pipeline.alignment = WD_ALIGN_PARAGRAPH.CENTER
pipeline.paragraph_format.space_before = Pt(8)
pipeline.paragraph_format.space_after = Pt(10)
for run in pipeline.runs:
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Consolas")
    run.font.size = Pt(9.5)

add_before(document, conclusion, "FastAPI validates the protected artifact bundle at startup and loads the HGNN once. It uses CUDA when available or CPU otherwise, sets evaluation mode and disables gradient recording. The endpoint applies sigmoid exactly once, rejects non-finite values, sorts by descending score with the event name as a deterministic tie-breaker and returns the configured top K labels.")

add_before(document, conclusion, "25.6 API Persistence and Failure Behavior", "Heading 2")
api_rows = [
    ("Interface", "Responsibility"),
    ("GET v1 hgnn status", "Reports readiness, versions, checkpoint hash, vocabulary, graph schema, device and audit state"),
    ("POST v1 hgnn predict events", "Validates supported inputs and returns actual ordered baseline event scores"),
    ("GET application adverse event risks", "Reuses a matching persisted event profile or generates one"),
    ("POST application adverse event risks", "Retries or regenerates the profile for the exact current input"),
    ("HgnnEventPrediction", "Stores versions, stable input hash, result, coverage, status and generation time"),
]
table = document.add_table(rows=len(api_rows), cols=2)
for row_index, values in enumerate(api_rows):
    for column_index, value in enumerate(values):
        table.cell(row_index, column_index).text = value
style_table(table, [2.15, 4.35])
conclusion._p.addprevious(table._element)
add_before(document, conclusion, "Unknown medicines or indications are identified and produce limited coverage. Missing files, hash or vocabulary mismatch, incompatible architecture, timeout, unreachable service, malformed response or inference failure produce an explicit unavailable state with an empty event list. The system never changes missing evidence into zero risk.")

add_before(document, conclusion, "25.7 Interface Design", "Heading 2")
add_before(document, conclusion, "The responsive page follows the existing navy, clinical blue and teal design system and supports dark mode. It presents consultation context, wizard navigation, Baseline HGNN and Audit Pending badges, graph-coverage status, an ordered event list, numeric model-score text, accessible progress indicators, a short interpretation guide and expandable provenance. The bars use neutral blue and teal rather than clinical severity colors because the outputs are not validated low, moderate or high risk categories. Loading, empty, limited-coverage and unavailable states remain visually distinct. Navigation to Clinical Safety, Adverse Risk Assessment and Recommendations remains available even when HGNN inference fails.")

add_before(document, conclusion, "25.8 Isolation and Remaining Audit", "Heading 2")
add_before(document, conclusion, "The event profile is display-only. Its result is not supplied to LightGBM, bootstrap or conformal calculations, DDInter analysis, DrugCentral analysis, safety gates, candidate generation, recommendation weights, sorting or tie-breaking. Before any future patient-probability claim, the project still requires calibration, temporal and external validation, label-quality review, subgroup and fairness analysis, clinical usability testing, monitoring and formal governance approval.")

for paragraph in document.paragraphs:
    if paragraph.text == "prisma/schema.prisma":
        insertion_anchor = paragraph
        for reference in (
            "docs/HGNN_EVENT_PROFILE.md",
            "docs/HGNN_IMPLEMENTATION_REPORT.md",
            "ml/src/vitanexus_ml/inference/hgnn_predictor.py",
            "server/services/hgnnEventProvider.js",
            "src/pages/AdverseEventRisks.jsx",
        ):
            add_before(document, insertion_anchor, reference, "List Bullet")
        break

document.save(DOCUMENT_PATH)
print(DOCUMENT_PATH)
